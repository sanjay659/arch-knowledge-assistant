"""
Generate RAG Complete Guide — Batch 3 (Topics 9-12)
Run: python docs/generate_guide_batch3.py
Output: docs/RAG_Guide_Batch3.docx
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_cell_shading(cell, color_hex):
    shading = cell._element.get_or_add_tcPr()
    elem = shading.makeelement(qn('w:shd'), {qn('w:fill'): color_hex, qn('w:val'): 'clear'})
    shading.append(elem)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
        set_cell_shading(cell, "1F4E79")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
            if r_idx % 2 == 0:
                set_cell_shading(cell, "D6E4F0")
    return table


def add_body(doc, text):
    return doc.add_paragraph(text)


def add_bullet(doc, text):
    return doc.add_paragraph(text, style='List Bullet')


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)
    p.paragraph_format.left_indent = Cm(1)
    shading = p._element.get_or_add_pPr()
    elem = shading.makeelement(qn('w:shd'), {qn('w:fill'): 'F2F2F2', qn('w:val'): 'clear'})
    shading.append(elem)
    return p


# ── Create Document ──
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
os.makedirs("docs", exist_ok=True)

# ═══════════════════════════════════════════════════════════════
# TOPIC 9: RETRIEVAL — THE CORE OF RAG
# ═══════════════════════════════════════════════════════════════

doc.add_heading('9. Retrieval — The Core of RAG', level=1)

doc.add_heading('9.1 Why Retrieval Quality = Answer Quality', level=2)
add_body(doc,
    'The retriever is the most important component in the entire RAG system. '
    'If the retriever finds the wrong chunks, the LLM generates a wrong answer — '
    'no matter how powerful the model is.')
add_body(doc, 'Think of it this way:')
add_bullet(doc, 'Perfect retrieval + weak LLM = decent answers (LLM just summarizes good content)')
add_bullet(doc, 'Bad retrieval + powerful LLM = wrong answers (LLM confidently summarizes wrong content)')
add_bullet(doc, 'The LLM can only work with what the retriever gives it')

doc.add_heading('9.2 Two Types of Filtering — Used Together', level=2)
add_body(doc,
    'Our retriever combines two complementary search strategies:')

add_table(doc,
    ["Strategy", "How It Works", "Analogy", "Strength"],
    [
        ["Metadata Filtering", "Exact match on structured fields: client=DataStories", "SQL WHERE clause", "Precise — eliminates wrong clients instantly"],
        ["Semantic Search", "Meaning match on text: 'hosting setup' finds 'infrastructure deployment'", "Google-like search by meaning", "Intelligent — finds relevant content even with different words"],
    ]
)

add_body(doc, 'Used together:')
add_code(doc,
    'Query: "How is DataStories hosted?"\n\n'
    'Step 1 — Metadata filter: client = "DataStories"\n'
    '  89 chunks --> 20 DataStories chunks (eliminated 69 irrelevant chunks instantly)\n\n'
    'Step 2 — Semantic search within filtered set:\n'
    '  20 DataStories chunks ranked by meaning similarity to "hosted"\n'
    '  Top 8 returned: Slide 9 (0.61), Slide 10 (0.53), Slide 4 (0.52), ...\n\n'
    'Without metadata filtering:\n'
    '  All 89 chunks compete --> MediFlow hosting chunks might outrank DataStories ones\n'
    '  --> CROSS-CLIENT CONTAMINATION in the answer')

add_body(doc,
    'This is why metadata tagging in the chunker was so critical. '
    'Without proper client and section_type tags, filtered retrieval is impossible.')

doc.add_heading('9.3 Intent Detection', level=2)
add_body(doc,
    'Different questions need different retrieval strategies. '
    'The retriever detects the query intent BEFORE searching.')

add_table(doc,
    ["Intent", "Keywords That Trigger", "Retrieval Strategy", "LLM Model (Future)"],
    [
        ["comparison", "compare, vs, versus, difference, what changed", "TWO separate searches: current_architecture + proposed_architecture", "GPT-4.1 Full (complex reasoning)"],
        ["explanation", "explain, describe, overview, how does, how is, what is", "Standard search with broad top_k", "GPT-4.1 Mini (good enough)"],
        ["specific_component", "gateway, database, auth, VPN, firewall, port, API", "Standard search — semantic precision matters", "GPT-4.1 Mini"],
        ["listing", "list, which clients, how many, show me all", "Cross-client search (no client filter)", "GPT-4.1 Mini"],
        ["general", "(none of above matched)", "Standard search", "GPT-4.1 Mini (default)"],
    ]
)

doc.add_heading('9.4 Comparison Retrieval — Special Handling', level=2)
add_body(doc,
    'When someone asks "Compare current vs proposed for RetailEdge", '
    'a single search would return a random mix of current and proposed chunks. '
    'The LLM would struggle to organize a structured comparison.')
add_body(doc, 'Instead, we do TWO targeted retrievals:')
add_code(doc,
    'Query: "Compare current vs proposed for RetailEdge"\n\n'
    'Search 1: client=RetailEdge, section_type=current_architecture\n'
    '  --> 3 chunks about the current state\n\n'
    'Search 2: client=RetailEdge, section_type=proposed_architecture\n'
    '  --> 3 chunks about the proposed state\n\n'
    'Both passed to LLM as separate sections:\n'
    '  "## CURRENT ARCHITECTURE\n'
    '   [chunks about current state]\n\n'
    '   ## PROPOSED ARCHITECTURE\n'
    '   [chunks about proposed state]"\n\n'
    '--> LLM produces a structured Current vs Proposed comparison')

doc.add_heading('9.5 Auto Client Detection', level=2)
add_body(doc,
    'Users naturally say "Explain DataStories architecture" rather than '
    'first selecting a client from a dropdown. The retriever auto-detects '
    'the client name from the query text.')
add_code(doc,
    'Algorithm:\n'
    '1. Get all known client names from the indexed collection\n'
    '   --> ["DataStories", "MediFlow", "RetailEdge", "TechNova"]\n\n'
    '2. Check if any client name appears in the query (case-insensitive)\n'
    '   "How is DataStories hosted?" --> found "DataStories"\n'
    '   "Which clients use AWS?"     --> no match --> search ALL clients\n\n'
    'Limitation: Simple substring matching.\n'
    'Phase 2 improvement: Use Named Entity Recognition or fuzzy matching.')

doc.add_heading('9.6 Relevance Threshold', level=2)
add_body(doc,
    'Not all retrieved chunks are useful. We apply a minimum relevance threshold '
    'of 0.35 — chunks scoring below this are dropped.')
add_table(doc,
    ["Score Range", "Meaning", "Action"],
    [
        ["0.6 - 1.0", "Highly relevant — strong semantic match", "✅ Include (top priority)"],
        ["0.4 - 0.6", "Moderately relevant — related topic", "✅ Include (supporting context)"],
        ["0.35 - 0.4", "Weakly relevant — marginal connection", "✅ Include (borderline)"],
        ["0.0 - 0.35", "Not relevant enough", "❌ Drop (would add noise)"],
    ]
)
add_body(doc,
    'The threshold is tunable. If too many good chunks are dropped, lower it to 0.30. '
    'If too much noise appears in answers, raise it to 0.40. '
    'Evaluation results guide this tuning.')

doc.add_heading('9.7 Retrieval Metrics', level=2)
add_table(doc,
    ["Metric", "Definition", "What It Tells You"],
    [
        ["Precision", "Of chunks retrieved, how many are actually relevant?", "High precision = less noise in LLM context"],
        ["Recall", "Of all relevant chunks, how many were retrieved?", "High recall = complete answers, nothing missed"],
        ["Top-K", "How many chunks to retrieve per query (we use 8)", "Higher K = better recall but lower precision"],
        ["MRR (Mean Reciprocal Rank)", "How high does the first relevant chunk rank?", "Higher MRR = right answer found quickly"],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 10: GENERATION & PROMPT ENGINEERING
# ═══════════════════════════════════════════════════════════════

doc.add_heading('10. Generation & Prompt Engineering', level=1)

doc.add_heading('10.1 What the Generator Does', level=2)
add_body(doc,
    'The generator takes retrieved chunks + user question and produces '
    'a human-readable answer with citations. This is where the LLM (GPT-4.1) is used.')
add_code(doc,
    'Input:\n'
    '  System prompt: "Answer ONLY from context, include citations..."\n'
    '  Context: [Chunk 1: Slide 9 content] [Chunk 2: Slide 10 content] ...\n'
    '  Question: "How is DataStories hosted?"\n\n'
    'Output:\n'
    '  "DataStories is hosted on AWS with account 556008695729.\n'
    '   The main website datastories.com is on AWS.\n'
    '   Two additional websites on Unbounce will be decommissioned.\n'
    '   [Source: blueprint.pptx, Slide 9] [Source: blueprint.pptx, Slide 10]"')

doc.add_heading('10.2 The System Prompt — Most Important Text in the System', level=2)
add_body(doc,
    'The system prompt defines HOW the LLM behaves. Every word is intentional. '
    'A bad prompt with good retrieval = bad answers. '
    'A good prompt with good retrieval = great answers.')
add_body(doc, 'Our system prompt enforces 6 rules:')

add_table(doc,
    ["Rule", "What It Does", "Why It Matters"],
    [
        ["1. Answer ONLY from context", "LLM cannot use training data", "Prevents hallucination — answers are grounded"],
        ["2. Include citations", "[Source: file.pptx, Slide X] after every claim", "Users can verify — trust and auditability"],
        ["3. Say 'I don't know'", "If context is insufficient, admit it", "Prevents confident wrong answers"],
        ["4. Structure clearly", "Use headings, bullets, numbered lists", "Readable, scannable answers"],
        ["5. Comparison format", "Separate Current vs Proposed sections", "Clean comparisons, not jumbled text"],
        ["6. Be precise", "Use exact numbers, names, terms from context", "Accurate, not paraphrased incorrectly"],
    ]
)

doc.add_heading('10.3 Grounded Generation — Why It Matters', level=2)
add_body(doc, 'The key difference between RAG and raw LLM:')
add_table(doc,
    ["Aspect", "Raw LLM", "RAG (Grounded)"],
    [
        ["Source of answer", "Training data (may be outdated)", "YOUR documents (always current)"],
        ["Hallucination risk", "High — model fills gaps with plausible fiction", "Low — model only uses provided context"],
        ["Verifiability", "Cannot verify — no source", "Every claim cites [Source: file, Slide X]"],
        ["Confidence calibration", "Often overconfident about wrong answers", "Says 'I don't have enough info' when context is insufficient"],
    ]
)

add_body(doc, 'Real example from our system:')
add_code(doc,
    'Query: "Compare current vs proposed for DataStories"\n\n'
    'Retrieved: 2 current_architecture chunks, 0 proposed_architecture chunks\n\n'
    'LLM response (grounded behavior):\n'
    '  "### Current Architecture\n'
    '   DataStories currently hosts servers in ACN approved Cloud platforms...\n'
    '   [Source: blueprint.pptx, Slide 8]\n\n'
    '   ### Proposed Architecture\n'
    '   No relevant documents found for the proposed architecture.\n\n'
    '   ### Key Differences\n'
    '   I don\'t have enough information in the available documents\n'
    '   to describe the proposed architecture or highlight differences."\n\n'
    'The LLM did NOT hallucinate a proposed architecture.\n'
    'It honestly said "I don\'t have enough information." THIS IS CORRECT BEHAVIOR.')

doc.add_heading('10.4 Temperature — Why We Use 0.1', level=2)
add_table(doc,
    ["Temperature", "Behavior", "Best For"],
    [
        ["0.0", "Completely deterministic — same input always gives same output", "Testing, evaluation"],
        ["0.1", "Almost deterministic with tiny variation", "Factual RAG answers (our choice)"],
        ["0.3-0.5", "Some creativity while staying mostly factual", "Creative writing with facts"],
        ["0.7-1.0", "High creativity, varied outputs", "Brainstorming, poetry, fiction"],
    ]
)
add_body(doc,
    'For architecture documentation, we want FACTUAL answers, not creative ones. '
    'Temperature 0.1 ensures consistent, reliable responses while allowing '
    'slight natural variation in phrasing.')

doc.add_heading('10.5 Context Building — How Chunks Become a Prompt', level=2)
add_body(doc,
    'The generator formats retrieved chunks into a numbered, labeled context block:')
add_code(doc,
    '## RETRIEVED CONTEXT\n\n'
    '[1] Source: blueprint.pptx, Slide 9\n'
    '    Client: DataStories | Section: current_architecture\n'
    '    Content:\n'
    '    Existing Hosting - Overview Internet\n'
    '    https://demo.datastories.com 669412377329...\n\n'
    '[2] Source: blueprint.pptx, Slide 10\n'
    '    Client: DataStories | Section: integration\n'
    '    Content:\n'
    '    DataStories has one AWS account...\n\n'
    '## Question\n'
    'How is DataStories currently hosted?')
add_body(doc, 'Why numbered?')
add_bullet(doc, 'Helps the LLM reference specific chunks in its answer')
add_bullet(doc, 'Helps humans debug: "Chunk [3] was irrelevant — need better filtering"')
add_body(doc, 'Why include metadata labels?')
add_bullet(doc, 'LLM needs source info to generate citations: [Source: blueprint.pptx, Slide 9]')
add_bullet(doc, 'Section type helps LLM understand context: this is about current state vs proposed state')

doc.add_heading('10.6 Conversation Memory', level=2)
add_body(doc,
    'Users often ask follow-up questions:')
add_code(doc,
    'Q1: "How is DataStories hosted?"\n'
    'A1: "DataStories is hosted on AWS with..."\n\n'
    'Q2: "What about the security standards?"\n'
    '    ^-- "the" refers to DataStories implicitly')
add_body(doc,
    'Without conversation memory, Q2 would search ALL clients because '
    '"DataStories" is not mentioned. With memory, the LLM sees the previous '
    'Q&A pair and understands the implicit reference.')
add_body(doc,
    'We keep a sliding window of the last 5 Q&A pairs. This balances context '
    'continuity with token cost (more history = more input tokens = higher cost).')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 11: LLM MODEL SELECTION FRAMEWORK
# ═══════════════════════════════════════════════════════════════

doc.add_heading('11. LLM Model Selection Framework', level=1)

doc.add_heading('11.1 Available Models on Azure OpenAI', level=2)
add_body(doc, 'The GPT-4.1 family offers three tiers for different needs and budgets:')

add_table(doc,
    ["Model", "Input $/1M tokens", "Output $/1M tokens", "Context Window", "Best For"],
    [
        ["GPT-4.1 (Full)", "$2.00", "$8.00", "1.05M tokens", "Complex reasoning, comparisons, highest quality"],
        ["GPT-4.1 Mini", "$0.40", "$1.60", "1.05M tokens", "Standard Q&A, explanations — best cost/quality balance"],
        ["GPT-4.1 Nano", "$0.10", "$0.40", "1.05M tokens", "Classification, simple extraction, high-volume tasks"],
    ]
)

add_body(doc, 'Key insight: Mini is 80% cheaper than Full but performance is roughly equivalent to GPT-4o on most tasks. The gap widens only on frontier benchmarks requiring deep reasoning.')

doc.add_heading('11.2 The 5 Decision Factors', level=2)
add_body(doc, 'Use these factors to choose the right model for any RAG system:')

add_table(doc,
    ["Factor", "Points to Nano", "Points to Mini", "Points to Full"],
    [
        ["1. Task Complexity", "Simple classification, extraction", "Standard Q&A, summaries, explanations", "Complex reasoning, multi-doc comparison"],
        ["2. Query Volume", "> 1000 queries/day", "100-1000 queries/day", "< 100 queries/day (cost irrelevant)"],
        ["3. Accuracy Requirement", "Nice to have (internal tool)", "Must be correct (team decisions)", "Lives depend on it (medical, legal)"],
        ["4. Output Complexity", "Short answers (1-2 sentences)", "Structured explanations with bullets", "Detailed comparisons, analysis, tables"],
        ["5. Monthly Budget", "Rs.100-200/month", "Rs.500-1000/month", "Rs.2000-5000/month"],
    ]
)

doc.add_heading('11.3 Our Multi-Model Architecture', level=2)
add_body(doc,
    'The smartest approach is NOT "use one model for everything." '
    'Instead, route different query types to different models:')
add_code(doc,
    'User Question\n'
    '    |\n'
    '    v\n'
    'Intent Detection\n'
    '    |\n'
    '    +-- classification task    --> GPT-4.1 Nano   ($0.0002/call)\n'
    '    +-- simple Q&A             --> GPT-4.1 Mini   ($0.0029/call)\n'
    '    +-- explanation            --> GPT-4.1 Mini   ($0.0029/call)\n'
    '    +-- comparison/reasoning   --> GPT-4.1 Full   ($0.0144/call)\n\n'
    'Average cost per query: ~$0.004 (Rs.0.38)\n'
    'Monthly (1,800 queries): ~$7.20 (Rs.683)')

add_body(doc, 'This is CHEAPER than using Full for everything AND BETTER quality where it matters.')

doc.add_heading('11.4 Cost Comparison for Our Scenario', level=2)
add_body(doc, 'Scenario: 6 architects, 10 queries/day each = 60 queries/day = 1,800/month')
add_body(doc, 'Per query: ~4,000 input tokens + ~800 output tokens')

add_table(doc,
    ["Strategy", "Per Query Cost", "Monthly USD", "Monthly INR"],
    [
        ["GPT-4.1 Full for ALL queries", "$0.0144", "$25.92", "Rs.2,460"],
        ["GPT-4.1 Mini for ALL queries", "$0.0029", "$5.18", "Rs.492"],
        ["GPT-4.1 Nano for ALL queries", "$0.00072", "$1.30", "Rs.123"],
        ["Multi-model (smart routing)", "~$0.004", "~$7.20", "~Rs.683"],
    ]
)

add_body(doc, 'Per query cost calculation example (GPT-4.1 Mini):')
add_code(doc,
    'Input:  4,000 tokens x ($0.40 / 1,000,000) = $0.0016\n'
    'Output:   800 tokens x ($1.60 / 1,000,000) = $0.00128\n'
    '                                       Total = $0.00288 per query\n\n'
    'Monthly: $0.00288 x 1,800 = $5.18\n'
    'INR:     $5.18 x 94.91 = Rs.492')

doc.add_heading('11.5 Decision Tree', level=2)
add_code(doc,
    'START\n'
    '  |\n'
    '  +-- "How many queries per day?"\n'
    '  |     +-- < 100/day    --> Cost is negligible, choose by QUALITY\n'
    '  |     +-- 100-1000/day --> Mini models are the sweet spot\n'
    '  |     +-- > 1000/day   --> Nano + optimize aggressively\n'
    '  |\n'
    '  +-- "How complex are the answers?"\n'
    '  |     +-- Simple (yes/no, short facts)       --> Nano\n'
    '  |     +-- Medium (explanations, summaries)    --> Mini\n'
    '  |     +-- Complex (comparisons, analysis)     --> Full\n'
    '  |\n'
    '  +-- "What is the cost of a WRONG answer?"\n'
    '  |     +-- Low (internal tool, can verify)     --> Mini or Nano\n'
    '  |     +-- Medium (team decisions based on it) --> Mini or Full\n'
    '  |     +-- High (customer-facing, compliance)  --> Full + guardrails\n'
    '  |\n'
    '  +-- "Can I test and switch later?"\n'
    '        +-- Yes --> Start with Mini, upgrade if needed\n'
    '        +-- No  --> Start with Full for safety')

doc.add_heading('11.6 Practical Recommendation', level=2)
add_body(doc, 'For our Architecture Knowledge Assistant:')
add_bullet(doc, 'START with GPT-4.1 Mini for everything (simplest, Rs.492/month)')
add_bullet(doc, 'RUN evaluation — check quality scores')
add_bullet(doc, 'IF comparison queries score low — upgrade THOSE to GPT-4.1 Full')
add_bullet(doc, 'IF everything scores high — stay with Mini (why pay more?)')
add_bullet(doc, 'TRACK token usage in API responses (already built into our GeneratedAnswer)')
add_body(doc, 'This is how production systems evolve: start simple, measure, then optimize.')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 12: EMBEDDING MODEL SELECTION
# ═══════════════════════════════════════════════════════════════

doc.add_heading('12. Embedding Model Selection', level=1)

doc.add_heading('12.1 Available Embedding Models', level=2)
add_table(doc,
    ["Model", "Dimensions", "Price / 1M tokens", "Quality (MIRACL)", "Best For"],
    [
        ["text-embedding-3-small", "1,536", "$0.02", "44.0", "Most use cases — best cost/quality ratio"],
        ["text-embedding-3-large", "3,072", "$0.13", "54.9", "Complex semantic tasks, multi-language"],
        ["text-embedding-ada-002", "1,536", "$0.10", "31.4", "Legacy — do not use (older, more expensive)"],
    ]
)

doc.add_heading('12.2 Why We Chose text-embedding-3-small', level=2)
add_bullet(doc, '6.5x cheaper than text-embedding-3-large ($0.02 vs $0.13 per 1M tokens)')
add_bullet(doc, 'Quality difference is marginal for English enterprise documents')
add_bullet(doc, '1,536 dimensions is more than sufficient for ~100 chunks')
add_bullet(doc, 'Same max input (8,191 tokens) as the large model')
add_bullet(doc, 'If evaluation shows retrieval quality issues, THEN try large — but not before measuring')

doc.add_heading('12.3 When Would You Need text-embedding-3-large?', level=2)
add_bullet(doc, 'Multi-language documents (e.g., Japanese + English architecture docs)')
add_bullet(doc, 'Very domain-specific technical jargon with subtle distinctions')
add_bullet(doc, 'Millions of chunks where precision improvements compound')
add_bullet(doc, 'Scientific, medical, or legal documents with nuanced terminology')
add_body(doc, 'For our architecture PPTs (English, enterprise, ~100 chunks)? Small is plenty.')

doc.add_heading('12.4 Embedding Cost Analysis', level=2)
add_body(doc, 'Embedding cost for our system:')
add_code(doc,
    'INDEXING (one-time per batch):\n'
    '  89 chunks x ~300 avg tokens = 26,700 tokens\n'
    '  26,700 x ($0.02 / 1,000,000) = $0.000534\n'
    '  = Rs.0.05 (five paise!) -- essentially FREE\n\n'
    'QUERYING (per question):\n'
    '  1 query x ~20 tokens = 20 tokens\n'
    '  20 x ($0.02 / 1,000,000) = $0.0000004\n'
    '  = Rs.0.00004 -- truly negligible\n\n'
    'MONTHLY (1,800 queries + occasional re-indexing):\n'
    '  Embedding cost: < Rs.1/month\n'
    '  This is so cheap that embedding model cost should NEVER\n'
    '  be a factor in your model selection decision.')

doc.add_heading('12.5 The Critical Rule: Same Model for Indexing AND Querying', level=2)
add_body(doc,
    'This is the most important rule for embeddings in RAG. Violating it '
    'produces zero results or garbage results.')
add_code(doc,
    'CORRECT:\n'
    '  Indexing:  text-embedding-3-small --> 1536 dimensions\n'
    '  Querying:  text-embedding-3-small --> 1536 dimensions\n'
    '  Compare:   1536 vs 1536 --> WORKS\n\n'
    'WRONG:\n'
    '  Indexing:  text-embedding-3-small --> 1536 dimensions\n'
    '  Querying:  text-embedding-3-large --> 3072 dimensions\n'
    '  Compare:   1536 vs 3072 --> CRASH (dimension mismatch)\n\n'
    'ALSO WRONG:\n'
    '  Indexing:  text-embedding-3-small --> 1536 dimensions\n'
    '  Querying:  ChromaDB default model --> 384 dimensions\n'
    '  Compare:   1536 vs 384 --> CRASH (this is the bug we actually hit!)')

add_body(doc, 'Consequences of changing the embedding model:')
add_table(doc,
    ["Change", "Impact", "Action Required"],
    [
        ["Changed LLM (e.g., Mini to Full)", "No impact on embeddings", "No re-indexing needed"],
        ["Changed chunk size", "Old chunks have wrong boundaries", "Must re-index ALL documents"],
        ["Changed embedding model", "Old vectors have wrong dimensions", "Must re-index ALL documents"],
        ["Added new documents", "New docs need indexing", "Run ingestion for new docs only"],
        ["Updated existing document", "Old chunks are stale", "Run ingestion — idempotent upsert handles it"],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# SAVE
# ═══════════════════════════════════════════════════════════════

output_path = "docs/RAG_Guide_Batch3.docx"
doc.save(output_path)
print(f"✅ Saved: {output_path}")
print(f"   Topics covered: 9-12")
print(f"   (Retrieval, Generation, LLM Selection, Embedding Selection)")