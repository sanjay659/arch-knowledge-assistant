"""
Generate RAG Complete Guide — Batch 1 (Topics 1-4)
Run: python docs/generate_guide_batch1.py
Output: docs/RAG_Guide_Batch1.docx
"""

import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# ── Helpers ──

def set_cell_shading(cell, color_hex):
    """Set cell background color."""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear',
    })
    shading.append(shading_elem)


def add_table(doc, headers, rows, col_widths=None):
    """Add a formatted table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
        set_cell_shading(cell, "1F4E79")

    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
            if r_idx % 2 == 0:
                set_cell_shading(cell, "D6E4F0")

    return table


def add_body(doc, text):
    """Add body paragraph."""
    p = doc.add_paragraph(text)
    p.style.font.size = Pt(11)
    return p


def add_bullet(doc, text, level=0):
    """Add bullet point."""
    p = doc.add_paragraph(text, style='List Bullet')
    if level > 0:
        p.style = doc.styles['List Bullet 2'] if 'List Bullet 2' in [s.name for s in doc.styles] else p.style
    return p


def add_code_block(doc, text):
    """Add monospaced text block."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)
    pf = p.paragraph_format
    pf.left_indent = Cm(1)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    # Light gray background via shading
    shading = p._element.get_or_add_pPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): 'F2F2F2',
        qn('w:val'): 'clear',
    })
    shading.append(shading_elem)
    return p


# ── Create Document ──

doc = Document()

# Set default font
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

os.makedirs("docs", exist_ok=True)

# ═══════════════════════════════════════════════════════════════
# TITLE PAGE
# ═══════════════════════════════════════════════════════════════

doc.add_paragraph()  # Spacer
doc.add_paragraph()
doc.add_paragraph()

title = doc.add_heading('The Complete RAG Guide', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_heading('From Theory to Production', level=1)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()

project = doc.add_paragraph('Architecture Knowledge Assistant')
project.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in project.runs:
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(50, 50, 50)

doc.add_paragraph()

tagline = doc.add_paragraph(
    'A hands-on guide covering every layer of RAG — '
    'from document extraction to LLM generation, '
    'with real code, real bugs, and real cost numbers.'
)
tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in tagline.runs:
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph()
doc.add_paragraph()

meta = doc.add_paragraph('Author: Architecture Team\nDate: June 2026\nVersion: 1.0')
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 1: WHAT IS RAG?
# ═══════════════════════════════════════════════════════════════

doc.add_heading('1. What is RAG?', level=1)

doc.add_heading('1.1 Definition', level=2)
add_body(doc,
    'RAG (Retrieval-Augmented Generation) is a pattern where an LLM generates answers '
    'using information RETRIEVED from your own documents, rather than relying solely on '
    'its training data.'
)
add_body(doc,
    'In simple terms: instead of asking the AI "what do you know?", you first FIND '
    'the relevant documents, then ask the AI "based on THESE documents, answer my question."'
)

doc.add_heading('1.2 How RAG Works (The 30-Second Version)', level=2)
add_body(doc, 'RAG has two phases:')
add_body(doc,
    'Phase 1 — INGESTION (one-time per document):\n'
    'Your PPT → Extract text → Split into chunks → Convert to vectors → Store in database'
)
add_body(doc,
    'Phase 2 — QUERY (every time someone asks a question):\n'
    'User question → Convert to vector → Find similar chunks → Pass chunks + question to LLM → Answer with citations'
)
add_body(doc,
    'The key insight: the LLM never sees your entire document collection. '
    'It only sees the 5-8 most relevant chunks for each specific question. '
    'This keeps answers focused, accurate, and affordable.'
)

doc.add_heading('1.3 RAG vs Alternatives', level=2)
add_body(doc,
    'When you want an LLM to know about YOUR data, you have three options:'
)

add_table(doc,
    ["Approach", "How It Works", "Pros", "Cons", "Cost", "Best For"],
    [
        ["Raw LLM", "Ask GPT directly, no custom data", "Simple, fast", "Hallucinations, no private data, no citations", "API cost only", "General knowledge Q&A"],
        ["Fine-tuning", "Retrain the model on your data", "Deep domain knowledge", "Very expensive ($1000s+), data goes stale, hard to update", "$$$$", "Massive stable datasets"],
        ["RAG", "Retrieve docs at query time, pass to LLM", "Fresh data, cited, affordable, updatable", "Needs pipeline engineering, retrieval quality matters", "$-$$", "Enterprise docs, knowledge bases"],
    ]
)

add_body(doc,
    'RAG is the sweet spot for most enterprise use cases because: your documents change frequently, '
    'answers must be traceable to source documents, you cannot afford to fine-tune every time a PPT '
    'is updated, and you need different clients\' data kept separate.'
)

doc.add_heading('1.4 The RAG Iceberg', level=2)
add_body(doc,
    'Most RAG tutorials only cover what\'s "above the water" — the basics that work for demos. '
    'Production RAG requires everything "below the water" — the hidden complexity that makes systems reliable.'
)

add_body(doc, 'ABOVE THE WATER (What beginners learn):')
add_bullet(doc, 'LangChain / LlamaIndex basics')
add_bullet(doc, 'Data loaders (PDF, PPT, CSV)')
add_bullet(doc, 'Chunking and embeddings')
add_bullet(doc, 'Vector databases (FAISS, Chroma, Pinecone)')
add_bullet(doc, 'Prompt templates')
add_bullet(doc, 'Basic retrieval + generation')

add_body(doc, 'BELOW THE WATER (What builders must handle):')
add_bullet(doc, 'Preprocessing & data cleaning')
add_bullet(doc, 'Section type classification')
add_bullet(doc, 'Metadata filtering')
add_bullet(doc, 'Query intent detection')
add_bullet(doc, 'Reranking (cross-encoders)')
add_bullet(doc, 'Evaluation metrics (retrieval hit, faithfulness, citation)')
add_bullet(doc, 'Hallucination control')
add_bullet(doc, 'Multi-hop retrieval')
add_bullet(doc, 'PII masking')
add_bullet(doc, 'Caching & latency optimization')
add_bullet(doc, 'Feedback loops & continuous improvement')
add_bullet(doc, 'LLM model routing (different models for different tasks)')
add_bullet(doc, 'Cost optimization')
add_bullet(doc, 'Secure retrieval (RBAC)')

add_body(doc,
    'Our project covers BOTH layers — that is what makes this guide different from typical tutorials.'
)

doc.add_heading('1.5 Why RAG is the Right Choice for Enterprise', level=2)
add_table(doc,
    ["Requirement", "Raw LLM", "Fine-tuning", "RAG"],
    [
        ["Uses private/internal documents", "❌ No", "⚠️ Baked into model", "✅ Retrieved at query time"],
        ["Data updates frequently", "N/A", "❌ Must retrain ($$$)", "✅ Re-index in minutes"],
        ["Answers must cite sources", "❌ Cannot cite", "❌ Cannot cite", "✅ [Source: file, Slide X]"],
        ["Multiple data sources/clients", "❌ No separation", "❌ One model for all", "✅ Metadata filtering"],
        ["Affordable at scale", "✅ Cheapest", "❌ Most expensive", "✅ Moderate"],
        ["Handles follow-up questions", "✅ With context", "✅ With context", "✅ Conversation memory"],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 2: OUR PROBLEM & SOLUTION
# ═══════════════════════════════════════════════════════════════

doc.add_heading('2. Our Problem & Solution', level=1)

doc.add_heading('2.1 The Problem: Knowledge Silos', level=2)
add_body(doc,
    'Our architecture team has 6 members, each working on different clients. '
    'Each architect maintains their own client-specific architecture documents '
    '(PowerPoint presentations, PDFs, design documents). This creates knowledge silos where:'
)
add_bullet(doc, 'No one has visibility into other clients\' architectures')
add_bullet(doc, 'Finding information requires manually opening and searching through PPT files')
add_bullet(doc, 'Onboarding new team members takes weeks of knowledge transfer')
add_bullet(doc, 'Cross-client insights (e.g., "which clients use AWS?") are nearly impossible')
add_bullet(doc, 'When an architect is unavailable, their client knowledge is inaccessible')

doc.add_heading('2.2 The Solution: Architecture Knowledge Assistant', level=2)
add_body(doc,
    'A RAG-powered system where all architects drop their client architecture PPTs into a '
    'shared folder structure, and anyone can ask natural language questions about any client\'s '
    'architecture and get clear, cited answers instantly.'
)
add_body(doc, 'How it works:')
add_bullet(doc, 'Each client has a folder: data/documents/ClientName/')
add_bullet(doc, 'Drop PPTs/PDFs into the client folder')
add_bullet(doc, 'System automatically indexes the content')
add_bullet(doc, 'Ask: "Explain DataStories architecture" → get a structured answer with citations')
add_bullet(doc, 'Ask: "Compare current vs proposed for RetailEdge" → get a side-by-side comparison')

doc.add_heading('2.3 Why This is a Good RAG Use Case', level=2)
add_table(doc,
    ["Criterion", "Our Scenario", "RAG Fit"],
    [
        ["Knowledge type", "Architecture PPTs (semi-structured text + tables)", "✅ Perfect for text extraction + retrieval"],
        ["Data freshness", "PPTs updated weekly/monthly", "✅ RAG re-indexes easily"],
        ["Question types", "Natural language (explain, compare, what is)", "✅ Exactly what RAG handles"],
        ["Answer source", "Must come from actual documents", "✅ RAG provides citations"],
        ["Data sensitivity", "Client-specific, internal only", "✅ Data stays local"],
        ["User count", "6 architects (low volume)", "✅ Cost-effective at this scale"],
        ["Multiple sources", "Different clients = different PPTs", "✅ Metadata filtering separates them"],
    ]
)

doc.add_heading('2.4 Test Dataset', level=2)
add_body(doc, 'We tested with 4 client architecture documents:')
add_table(doc,
    ["Client", "Document", "Slides", "Industry", "Migration Type"],
    [
        ["DataStories", "Hosting Integration Blueprint", "26", "Analytics/SaaS", "AWS → Accenture GMCS"],
        ["TechNova", "Hosting Blueprint", "17", "SaaS Platform", "AWS → Azure"],
        ["MediFlow", "Cloud Migration Blueprint", "14", "Healthcare", "On-Prem → AWS"],
        ["RetailEdge", "Architecture Blueprint", "15", "E-Commerce/Retail", "Hybrid Multi-Cloud"],
    ]
)
add_body(doc, 'Total: 72 slides → 67 meaningful slides → 89 indexed chunks → searchable in under 1 second.')

doc.add_heading('2.5 Question Types Supported', level=2)
add_table(doc,
    ["#", "Question Type", "Example", "Retrieval Strategy"],
    [
        ["1", "Explain architecture", "Explain DataStories architecture", "Multi-chunk summary, client filtered"],
        ["2", "Current design", "How is MediFlow currently hosted?", "Filter section_type = current_architecture"],
        ["3", "Proposed design", "What is the target architecture?", "Filter section_type = proposed_architecture"],
        ["4", "Compare", "Compare current vs proposed for RetailEdge", "Two separate retrievals → structured comparison"],
        ["5", "Budget", "What is the budget for TechNova?", "Retrieve budget-related chunks"],
        ["6", "Security", "What security standards apply?", "Retrieve security/compliance chunks"],
        ["7", "Cross-client", "Which clients use AWS?", "Search ALL clients (no filter)"],
        ["8", "Specific component", "How does the API gateway work?", "Precise semantic search"],
        ["9", "Follow-up", "What about the risks?", "Uses conversation memory from previous Q&A"],
        ["10", "Listing", "List all migration milestones", "Retrieve timeline/milestone chunks"],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 3: ARCHITECTURE OVERVIEW
# ═══════════════════════════════════════════════════════════════

doc.add_heading('3. Architecture Overview', level=1)

doc.add_heading('3.1 Phase 1 — Local Self-Hosted (Current)', level=2)
add_body(doc,
    'Phase 1 is designed for learning and proof-of-concept. '
    'Everything runs on a single machine (or VM) with minimal cloud dependencies.'
)

add_body(doc, 'Components:')
add_table(doc,
    ["Component", "Technology", "Purpose"],
    [
        ["Document Store", "Local shared folder", "Client PPTs organized by folder"],
        ["Text Extraction", "python-pptx, PyMuPDF, python-docx", "Extract text from PPT/PDF/DOCX"],
        ["Preprocessing", "Custom Python (regex + keywords)", "Clean boilerplate, classify sections"],
        ["Chunking", "Custom Python", "Split into 100-1500 char chunks with metadata"],
        ["Embeddings", "Azure OpenAI text-embedding-3-small", "Convert text → 1536-dim vectors"],
        ["Vector Store", "ChromaDB (local)", "Store and search vectors"],
        ["LLM Generation", "Azure OpenAI GPT-4.1", "Generate answers from retrieved chunks"],
        ["API", "FastAPI (port 8000)", "REST endpoints for query and ingestion"],
        ["UI", "Streamlit (port 8501)", "Chat interface for the team"],
    ]
)

add_body(doc, 'Estimated monthly cost: ₹8,000 – ₹11,000 (VM + AI tokens)')

doc.add_heading('3.2 Phase 2 — Azure-Native Managed (Future)', level=2)
add_body(doc,
    'Phase 2 replaces self-hosted components with Azure managed services '
    'for better quality, security, and scalability. The core RAG logic stays the same — '
    'only the infrastructure changes.'
)

add_body(doc, 'Estimated monthly cost: ₹15,000 – ₹18,000')

doc.add_heading('3.3 Component Migration Map', level=2)
add_body(doc, 'Each Phase 1 component maps directly to a Phase 2 Azure service:')

add_table(doc,
    ["Component", "Phase 1 (Current)", "Phase 2 (Future)", "Why Migrate"],
    [
        ["Document Store", "Local shared folder", "Azure Blob Storage", "Scalable, versioned, event triggers"],
        ["Text Extraction", "python-pptx / PyMuPDF", "Azure Document Intelligence", "OCR for diagrams, table structure, layout detection"],
        ["Vector Store", "ChromaDB", "Azure AI Search", "Hybrid search (vector + keyword), semantic reranking"],
        ["Search Type", "Semantic only", "Hybrid + Reranking", "Much better retrieval precision and recall"],
        ["Ingestion", "Manual Python script", "Azure Functions", "Auto-trigger on file upload (event-driven)"],
        ["API Host", "FastAPI on VM", "Azure App Service", "Managed, SSL, auto-scaling, authentication"],
        ["UI Host", "Streamlit on VM", "Azure App Service", "Professional, shareable URL"],
        ["Secrets", ".env file", "Azure Key Vault", "Secure, audited, rotatable keys"],
        ["Security", "None", "Azure RBAC + Entra ID", "Role-based document access per client"],
        ["Monitoring", "Console logs", "Azure Monitor + App Insights", "Dashboards, alerts, distributed tracing"],
    ]
)

add_body(doc,
    'Key design principle: The project is built MODULARLY so that each component can be '
    'replaced independently. Migrating from ChromaDB to Azure AI Search requires changing '
    'only 2 files (indexer.py and retriever.py). Everything else stays the same.'
)

doc.add_heading('3.4 Data Flow', level=2)
add_body(doc, 'WRITE PATH (Ingestion — runs once per document):')
add_code_block(doc,
    'PPT File\n'
    '  → Scanner (discover files, detect client from folder name)\n'
    '  → Extractor (PPT → text per slide, tables → pipe-delimited text)\n'
    '  → Preprocessor (remove boilerplate, classify section type)\n'
    '  → Chunker (merge small slides, split large ones, attach metadata)\n'
    '  → Indexer (embed via Azure OpenAI → store in ChromaDB)'
)

add_body(doc, 'READ PATH (Query — runs per question):')
add_code_block(doc,
    'User Question: "How is DataStories hosted?"\n'
    '  → Intent Detection → "explanation"\n'
    '  → Client Detection → "DataStories"\n'
    '  → Retriever (embed query → search ChromaDB with client filter → top 8 chunks)\n'
    '  → Generator (system prompt + chunks + question → GPT-4.1 → answer + citations)\n'
    '  → Response: structured answer with [Source: blueprint.pptx, Slide 9]'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 4: THE RAG PIPELINE EXPLAINED
# ═══════════════════════════════════════════════════════════════

doc.add_heading('4. The RAG Pipeline Explained', level=1)

doc.add_heading('4.1 Two Paths', level=2)
add_body(doc,
    'Every RAG system has two paths: the WRITE PATH (getting data in) and '
    'the READ PATH (getting answers out). Understanding both is essential.'
)

doc.add_heading('4.2 Write Path — Ingestion Pipeline', level=2)
add_body(doc,
    'The write path runs once per document (or when a document is updated). '
    'It transforms raw PPT files into searchable vector entries.'
)

add_table(doc,
    ["Step", "Module", "Input", "Output", "What Can Go Wrong"],
    [
        ["1. Scan", "scanner.py", "Folder path", "List of files + client names", "Missing files, wrong permissions, unsupported formats"],
        ["2. Extract", "extractor.py", "PPT/PDF file", "Text per slide/page + tables", "Images not extractable, broken shapes, encoding issues"],
        ["3. Preprocess", "preprocessor.py", "Raw text", "Clean text + section type label", "Boilerplate not caught, wrong section classification"],
        ["4. Chunk", "chunker.py", "Clean slides", "Sized chunks with metadata", "Wrong chunk size (too small/large), lost context at boundaries"],
        ["5. Embed", "indexer.py", "Chunk text", "1536-dimensional vector", "API rate limits, network errors, wrong model"],
        ["6. Store", "indexer.py", "Vector + metadata", "ChromaDB entry", "Duplicates if IDs not deterministic, disk space"],
    ]
)

add_body(doc, 'Key design decisions in the write path:')
add_bullet(doc, 'Folder name = client name: No manual labeling needed. Just create a folder called "DataStories" and drop files in it.')
add_bullet(doc, 'Slide-level extraction: Each PPT slide is extracted separately, preserving the architect\'s intended grouping.')
add_bullet(doc, 'Deterministic chunk IDs: Same file + same content = same ID. This makes re-indexing idempotent (safe to run multiple times).')
add_bullet(doc, 'Batch embedding: Chunks are embedded 16 at a time to avoid rate limits and reduce API overhead.')

doc.add_heading('4.3 Read Path — Query Pipeline', level=2)
add_body(doc,
    'The read path runs every time someone asks a question. '
    'It finds the most relevant chunks and generates a cited answer.'
)

add_table(doc,
    ["Step", "Module", "Input", "Output", "What Can Go Wrong"],
    [
        ["1. Intent Detection", "retriever.py", "Question text", "comparison / explanation / specific / listing / general", "Wrong intent → wrong retrieval strategy"],
        ["2. Client Detection", "retriever.py", "Question text", "Client name or None", "Client not detected → searches all clients (broader but noisier)"],
        ["3. Query Embedding", "retriever.py", "Question text", "1536-dim vector", "Must use SAME model as indexing (dimension mismatch otherwise)"],
        ["4. Retrieval", "retriever.py", "Query vector + filters", "Top-K relevant chunks with scores", "Wrong chunks retrieved, threshold too high/low"],
        ["5. Generation", "generator.py", "Chunks + question + system prompt", "Answer with citations", "Hallucination, missing citations, too verbose"],
        ["6. Response", "query_pipeline.py", "Generated answer + metadata", "QueryResponse object", "Timeout, token limit exceeded"],
    ]
)

add_body(doc, 'Key design decisions in the read path:')
add_bullet(doc, 'Intent-based routing: Comparison queries trigger TWO separate retrievals (current + proposed). All other queries use standard retrieval.')
add_bullet(doc, 'Auto client detection: Users naturally say "Explain DataStories architecture" — the system detects "DataStories" and filters automatically.')
add_bullet(doc, 'Relevance threshold: Chunks scoring below 0.35 similarity are dropped. This prevents low-quality context from confusing the LLM.')
add_bullet(doc, 'Conversation memory: Last 5 Q&A pairs are kept in memory, enabling follow-up questions like "What about the security?" after asking about a specific client.')

doc.add_heading('4.4 How the Paths Connect', level=2)
add_body(doc,
    'The write path creates the knowledge base. The read path queries it. '
    'They share the same embedding model (text-embedding-3-small) and the same '
    'vector store (ChromaDB). The critical rule: any change to the write path '
    '(e.g., different chunk size, different embedding model) requires RE-INDEXING '
    'all documents before the read path can work correctly.'
)

add_body(doc, 'Common scenarios:')
add_table(doc,
    ["Scenario", "What to Do"],
    [
        ["New PPT added", "Run ingestion → new chunks added to existing index"],
        ["PPT updated", "Run ingestion → old chunks deleted, new chunks added (idempotent)"],
        ["PPT deleted", "Delete chunks by source file, or re-index everything"],
        ["Changed chunk size", "Must re-index ALL documents (old chunks have wrong boundaries)"],
        ["Changed embedding model", "Must re-index ALL documents (old vectors have wrong dimensions)"],
        ["Changed LLM model", "No re-indexing needed (only affects generation, not storage)"],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# SAVE
# ═══════════════════════════════════════════════════════════════

output_path = "docs/RAG_Guide_Batch1.docx"
doc.save(output_path)
print(f"✅ Saved: {output_path}")
print(f"   Topics covered: 1-4")
print(f"   (What is RAG, Problem Statement, Architecture, Pipeline)")