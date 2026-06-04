"""
Generate RAG Complete Guide — Batch 5 (Topics 17-22) FINAL
Run: python docs/generate_guide_batch5.py
Output: docs/RAG_Guide_Batch5.docx
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_cell_shading(cell, color_hex):
    shading = cell._element.get_or_add_tcPr()
    elem = shading.makeelement(qn('w:shd'), {qn('w:fill'): color_hex, qn('w:val'): 'clear'})
    shading.append(elem)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
        set_cell_shading(cell, "1F4E79")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
            if r_idx % 2 == 0:
                set_cell_shading(cell, "D6E4F0")
    return table


def add_body(doc, text):
    return doc.add_paragraph(text)


def add_bullet(doc, text):
    return doc.add_paragraph(text, style='List Bullet')


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)
    p.paragraph_format.left_indent = Cm(1)
    shading = p._element.get_or_add_pPr()
    elem = shading.makeelement(qn('w:shd'), {qn('w:fill'): 'F2F2F2', qn('w:val'): 'clear'})
    shading.append(elem)
    return p


# ── Create Document ──
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
os.makedirs("docs", exist_ok=True)

# ═══════════════════════════════════════════════════════════════
# TOPIC 17: PHASE 2 MODERNIZATION PATH
# ═══════════════════════════════════════════════════════════════

doc.add_heading('17. Phase 2 Modernization Path', level=1)

doc.add_heading('17.1 Why Modernize?', level=2)
add_body(doc,
    'Phase 1 proves the concept and teaches the fundamentals. Phase 2 makes it '
    'enterprise-ready. The modernization is NOT about rewriting — it is about '
    'replacing self-hosted components with managed Azure services, one at a time.')
add_body(doc, 'Key benefits of Phase 2:')
add_bullet(doc, 'Better retrieval quality: Hybrid search (vector + keyword) catches what pure semantic search misses')
add_bullet(doc, 'Better extraction: Azure Document Intelligence reads text from architecture diagrams (OCR)')
add_bullet(doc, 'Better security: Key Vault for secrets, RBAC for per-client access control')
add_bullet(doc, 'Better observability: Azure Monitor dashboards, alerts, distributed tracing')
add_bullet(doc, 'Auto-scaling: App Service handles traffic spikes without manual VM management')
add_bullet(doc, 'Event-driven ingestion: Azure Functions auto-trigger when new PPTs are uploaded')

doc.add_heading('17.2 Component-by-Component Migration', level=2)

add_table(doc,
    ["Component", "Phase 1", "Phase 2", "Migration Effort", "Quality Impact"],
    [
        ["Vector Store", "ChromaDB (local)", "Azure AI Search", "Medium (change indexer + retriever)", "HIGH — hybrid search dramatically improves retrieval"],
        ["Text Extraction", "python-pptx, PyMuPDF", "Azure Document Intelligence", "Medium (change extractor)", "HIGH — OCR reads diagram text"],
        ["Ingestion Trigger", "Manual script", "Azure Functions (Blob trigger)", "Low (new function, same pipeline)", "Medium — auto-ingestion on upload"],
        ["API Hosting", "FastAPI on VM", "Azure App Service", "Low (deploy same code)", "Low — same functionality, managed hosting"],
        ["UI Hosting", "Streamlit on VM", "Azure App Service", "Low (deploy same code)", "Low — professional URL, SSL"],
        ["Secrets", ".env file", "Azure Key Vault", "Low (change config loading)", "Medium — secure, audited, rotatable"],
        ["Security", "None", "Azure RBAC + Entra ID", "High (new access control layer)", "HIGH — per-client document access"],
        ["Monitoring", "Console logs", "Azure Monitor + App Insights", "Medium (add SDK + configure)", "HIGH — dashboards, alerts, traces"],
    ]
)

doc.add_heading('17.3 Azure AI Search — The Biggest Quality Jump', level=2)
add_body(doc,
    'The single most impactful migration is ChromaDB to Azure AI Search. Here is why:')

add_table(doc,
    ["Feature", "ChromaDB (Phase 1)", "Azure AI Search (Phase 2)"],
    [
        ["Search type", "Semantic only (vector similarity)", "Hybrid: vector + BM25 keyword + semantic reranking"],
        ["Reranking", "None", "Built-in cross-encoder reranking"],
        ["Keyword search", "Not available", "BM25 full-text search (catches exact terms like 'port 443')"],
        ["Scoring", "Cosine distance only", "Combined score from vector + keyword + reranker"],
        ["Filtering", "Basic metadata filter", "OData filters with complex conditions"],
        ["Scaling", "Single machine", "Replicas + partitions for high availability"],
        ["Managed", "Self-hosted (you manage)", "Fully managed Azure service"],
    ]
)

add_body(doc, 'Why hybrid search matters:')
add_code(doc,
    'SEMANTIC ONLY (Phase 1):\n'
    '  Query: "port 443"\n'
    '  Embedding captures meaning: "network port, HTTPS"\n'
    '  May find: general networking chunks (good)\n'
    '  May miss: exact slide that mentions "443" specifically (bad)\n\n'
    'HYBRID (Phase 2):\n'
    '  Vector search: finds networking-related chunks\n'
    '  BM25 keyword search: finds chunks containing exact text "443"\n'
    '  Combined: finds BOTH semantic matches AND exact matches\n'
    '  Reranker: re-scores combined results for best final ranking\n\n'
    'Result: significantly better retrieval precision and recall')

doc.add_heading('17.4 Azure Document Intelligence — Reading Diagrams', level=2)
add_body(doc,
    'In Phase 1, architecture diagrams in PPTs are INVISIBLE to the system — '
    'they are images, and python-pptx cannot read text from images.')
add_body(doc, 'Azure Document Intelligence changes this:')
add_bullet(doc, 'Layout API: Detects document structure (headers, sections, tables, paragraphs)')
add_bullet(doc, 'Read API: OCR that extracts text from ANY image, including architecture diagrams')
add_bullet(doc, 'Table API: Properly extracts table structure including merged cells')
add_bullet(doc, 'Supports PPT, PDF, DOCX, images, and even scanned documents')
add_body(doc,
    'For architecture documents specifically, this means the system can finally '
    '"see" flow diagrams, network topology images, and integration diagrams that '
    'contain critical architectural information.')

doc.add_heading('17.5 Migration Strategy', level=2)
add_body(doc, 'Migrate incrementally, not all at once:')
add_table(doc,
    ["Week", "Migration", "Validates"],
    [
        ["Week 1-2", "Azure AI Search (replace ChromaDB)", "Hybrid search improves retrieval quality"],
        ["Week 3-4", "Azure Document Intelligence (replace python-pptx)", "OCR improves extraction from diagrams"],
        ["Week 5", "Azure App Service (deploy API + UI)", "Managed hosting, SSL, shared URL"],
        ["Week 6", "Azure Functions (event-driven ingestion)", "Auto-index on file upload"],
        ["Week 7", "Azure Key Vault + Monitor", "Security + observability"],
        ["Week 8", "RBAC + Entra ID", "Per-client access control"],
    ]
)
add_body(doc,
    'Each migration step can be evaluated independently. If Azure AI Search does not improve quality '
    '(unlikely), you can revert to ChromaDB. The modular design makes this possible.')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 18: MISSING LINKS & GAPS
# ═══════════════════════════════════════════════════════════════

doc.add_heading('18. Missing Links & Gaps to Address', level=1)

add_body(doc,
    'Our Phase 1 system works well for the POC, but there are gaps that a production '
    'enterprise system would need to address. These are listed in priority order.')

add_table(doc,
    ["#", "Gap", "What It Is", "Impact If Missing", "Effort", "When to Address"],
    [
        ["1", "Reranking", "Cross-encoder model re-scores retrieved chunks for better precision", "Some irrelevant chunks may appear in top results", "Medium", "Phase 2 (comes free with Azure AI Search)"],
        ["2", "Streaming Responses", "Token-by-token delivery to UI instead of waiting for full answer", "User waits 5-13 seconds seeing nothing, then full answer appears", "Low", "Phase 1.5 (can add to FastAPI + Streamlit)"],
        ["3", "Response Caching", "Cache frequent queries to avoid repeated LLM calls", "Same question costs the same every time, no latency benefit", "Low", "Phase 1.5 (Redis or in-memory dict)"],
        ["4", "PII Masking", "Detect and mask personal data before indexing", "Email addresses, phone numbers stored in vector DB", "Medium", "Phase 2 (use Azure AI Language PII detection)"],
        ["5", "Hallucination Detection", "Automated check if answer is grounded in retrieved context", "May pass hallucinated answers to users without flagging", "High", "Phase 2 (LLM-as-judge pattern)"],
        ["6", "Feedback Loop", "Users rate answers (thumbs up/down), system learns from feedback", "No way to know which answers are wrong without manual review", "Medium", "Phase 2"],
        ["7", "Multi-hop Retrieval", "Answer questions that require reasoning across multiple documents", "Cannot answer: 'Which clients with budget >$50K use AWS?'", "High", "Phase 3"],
        ["8", "Version-aware Retrieval", "Track document versions, answer based on latest or specific version", "May retrieve outdated information if PPT was updated", "Medium", "Phase 2"],
        ["9", "Role-based Access (RBAC)", "Users see only their authorized clients' data", "All users can query all clients (fine for our team, not for larger orgs)", "High", "Phase 2 (Azure Entra ID + search filters)"],
        ["10", "Observability Dashboard", "Track query volume, latency, error rates, token costs over time", "Cannot monitor system health or optimize costs", "Medium", "Phase 2 (Azure Monitor + App Insights)"],
    ]
)

doc.add_heading('18.1 Quick Wins (Can Be Added in Phase 1)', level=2)
add_body(doc, 'These gaps can be addressed with minimal effort in the current architecture:')

add_body(doc, 'Streaming Responses:')
add_code(doc,
    '# In generator.py — change create() to stream=True:\n'
    'response = openai_client.chat.completions.create(\n'
    '    model=deployment,\n'
    '    messages=messages,\n'
    '    stream=True,   # <-- add this\n'
    ')\n'
    'for chunk in response:\n'
    '    yield chunk.choices[0].delta.content  # Token by token')

add_body(doc, 'Response Caching:')
add_code(doc,
    '# Simple in-memory cache (add to query_pipeline.py):\n'
    'from functools import lru_cache\n\n'
    '@lru_cache(maxsize=100)\n'
    'def cached_query(question, client):\n'
    '    return self._actual_query(question, client)\n\n'
    '# Or use Redis for persistent caching across restarts')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 19: ADVANCED RAG PATTERNS
# ═══════════════════════════════════════════════════════════════

doc.add_heading('19. Advanced RAG Patterns (Future)', level=1)

add_body(doc,
    'Beyond Phase 2, there are advanced RAG architectures emerging in the industry. '
    'These are not needed for our current use case but represent where the field is heading.')

doc.add_heading('19.1 Agentic RAG', level=2)
add_body(doc, 'What: An AI agent that DECIDES how to answer, rather than following a fixed pipeline.')
add_code(doc,
    'User: "Which clients with budget over $50K are migrating to Azure?"\n\n'
    'Agent thinks:\n'
    '  1. I need to search for budget information across all clients\n'
    '  2. I need to filter for budgets > $50K\n'
    '  3. I need to check which of those are migrating to Azure\n'
    '  4. This requires multiple retrieval calls and reasoning\n\n'
    'Agent executes:\n'
    '  → Retrieve budget chunks for all clients\n'
    '  → Extract dollar amounts using a tool\n'
    '  → Filter: TechNova ($23K one-time + $13K/month), RetailEdge ($358K one-time)\n'
    '  → Check migration target for each\n'
    '  → Answer: "TechNova (migrating AWS to Azure) and RetailEdge (multi-cloud including Azure)"')
add_body(doc, 'When needed: Complex multi-step queries that require planning and tool use.')

doc.add_heading('19.2 Multimodal RAG', level=2)
add_body(doc, 'What: RAG that understands images, not just text.')
add_code(doc,
    'Current (text only):\n'
    '  Architecture diagram on Slide 9 → INVISIBLE (just an image)\n'
    '  "Explain the network diagram" → "I don\'t have enough information"\n\n'
    'Multimodal RAG:\n'
    '  Architecture diagram on Slide 9 → Vision model extracts:\n'
    '    "Internet → API Gateway → Auth Service → Backend → Database"\n'
    '  This text is indexed alongside regular slide text\n'
    '  "Explain the network diagram" → full answer from diagram content')
add_body(doc, 'When needed: When architecture diagrams contain critical information not available in text.')
add_body(doc, 'How: GPT-4o (vision) or Azure Document Intelligence Read API for image-to-text.')

doc.add_heading('19.3 Graph RAG', level=2)
add_body(doc, 'What: Combines vector search with a knowledge graph for relationship-aware retrieval.')
add_code(doc,
    'Traditional RAG:\n'
    '  "Which systems depend on the API Gateway?"\n'
    '  → Searches for chunks mentioning "API Gateway"\n'
    '  → May find some, miss others (depends on how text is written)\n\n'
    'Graph RAG:\n'
    '  Documents → Knowledge graph:\n'
    '    [API Gateway] --connects_to--> [Auth Service]\n'
    '    [API Gateway] --connects_to--> [Backend Cluster]\n'
    '    [Auth Service] --depends_on--> [Identity Provider]\n'
    '    [Backend Cluster] --uses--> [Database]\n\n'
    '  "Which systems depend on the API Gateway?"\n'
    '  → Graph traversal: API Gateway → Auth Service, Backend Cluster\n'
    '  → Answer includes ALL connected systems, not just those mentioned together')
add_body(doc, 'When needed: Complex relationship queries across many interconnected systems.')

doc.add_heading('19.4 Corrective RAG (CRAG)', level=2)
add_body(doc, 'What: Self-correcting RAG that checks and fixes its own answers.')
add_code(doc,
    'Standard RAG:\n'
    '  Retrieve → Generate → Return (hope it is correct)\n\n'
    'Corrective RAG:\n'
    '  Retrieve → Generate → CHECK if grounded → \n'
    '    If grounded → Return ✅\n'
    '    If NOT grounded → Re-retrieve with different strategy → Re-generate → Return')
add_body(doc, 'When needed: High-stakes environments where hallucination is unacceptable.')

doc.add_heading('19.5 Adaptive RAG', level=2)
add_body(doc, 'What: RAG that adjusts its strategy based on query complexity.')
add_code(doc,
    'Simple query: "What is DataStories?"\n'
    '  → Direct retrieval, fast answer, Nano model\n\n'
    'Medium query: "Explain the hosting architecture"\n'
    '  → Standard retrieval, detailed answer, Mini model\n\n'
    'Complex query: "Compare all clients\' migration approaches and recommend best practices"\n'
    '  → Multi-step retrieval, cross-client analysis, Full model + agent reasoning')
add_body(doc,
    'Our multi-model routing (Nano/Mini/Full per intent) is actually a simple form of Adaptive RAG. '
    'Full Adaptive RAG would also adjust chunk count, retrieval strategy, and post-processing per query.')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 20: NEXT GOALS & ROADMAP
# ═══════════════════════════════════════════════════════════════

doc.add_heading('20. Next Goals & Roadmap', level=1)

doc.add_heading('20.1 Short Term (1-2 Weeks)', level=2)
add_table(doc,
    ["Task", "Purpose", "Effort"],
    [
        ["Run full evaluation suite", "Get baseline quality numbers", "1 hour"],
        ["Tune chunk size based on eval results", "Improve retrieval precision", "2-3 hours"],
        ["Add multi-model routing (Nano/Mini/Full per intent)", "Optimize cost without sacrificing quality", "Half day"],
        ["Fix comparison retrieval for DataStories", "Lower threshold or remove section_type filter for comparison", "1 hour"],
        ["Add 5 more golden Q&A pairs per client", "Better evaluation coverage", "2 hours"],
        ["Add streaming responses to UI", "Better user experience (no waiting for full answer)", "Half day"],
    ]
)

doc.add_heading('20.2 Medium Term (1 Month)', level=2)
add_table(doc,
    ["Task", "Purpose", "Effort"],
    [
        ["Add response caching (Redis or in-memory)", "Reduce cost and latency for repeated queries", "1 day"],
        ["Add reranking for better retrieval precision", "Improve answer quality on ambiguous queries", "2 days"],
        ["Onboard 2-3 more real client PPTs", "Test system with more diverse documents", "1 day"],
        ["Create team demo presentation (PPT from this guide)", "Share learning with broader team", "2 days"],
        ["Add PII masking for email addresses in documents", "Privacy compliance", "1 day"],
        ["Set up automated evaluation in CI/CD", "Quality gate on every code change", "1 day"],
    ]
)

doc.add_heading('20.3 Long Term (3 Months)', level=2)
add_table(doc,
    ["Task", "Purpose", "Effort"],
    [
        ["Migrate to Azure AI Search (replace ChromaDB)", "Hybrid search + semantic reranking", "1 week"],
        ["Add Azure Document Intelligence (replace python-pptx)", "OCR for architecture diagrams", "1 week"],
        ["Deploy to Azure App Service", "Managed hosting with SSL and scaling", "2 days"],
        ["Implement RBAC with Azure Entra ID", "Per-client document access control", "1 week"],
        ["Add feedback loop (thumbs up/down on answers)", "Continuous improvement from user feedback", "3 days"],
        ["Add Azure Monitor + Application Insights", "Production observability dashboards", "2 days"],
        ["Implement hallucination detection (LLM-as-judge)", "Automated answer quality checking", "3 days"],
        ["Explore multimodal RAG for diagram understanding", "Read text from architecture diagrams", "1 week"],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 21: KEY TAKEAWAYS
# ═══════════════════════════════════════════════════════════════

doc.add_heading('21. Key Takeaways — 10 Lessons from Building This System', level=1)

add_body(doc,
    'These are the most important lessons from building the Architecture Knowledge Assistant. '
    'They apply to ANY RAG system, not just ours.')

add_table(doc,
    ["#", "Lesson", "Why It Matters"],
    [
        ["1", "Data quality > Model quality", "Cleaning boilerplate from PPTs improved answers more than any model upgrade would. Garbage in = garbage out, regardless of how powerful GPT is."],
        ["2", "Chunk size is the #1 tuning parameter", "Too small = fragmented context. Too large = diluted embeddings. Our sweet spot: 500-1500 chars (125-375 tokens). Always tune this first."],
        ["3", "Metadata filtering is as important as semantic search", "Without client filtering, answers mix different clients. With it, answers are precise. Invest in rich metadata during ingestion."],
        ["4", "Same embedding model for indexing AND querying", "Dimension mismatch = zero results. This is the most common RAG bug. We hit it ourselves (1536 vs 384)."],
        ["5", "Real documents break library assumptions", "python-pptx crashed on real PPTs. Always wrap extraction in try/except. One bad shape should not skip the entire slide."],
        ["6", "The system prompt is your most important prompt", "It enforces grounding ('only from context'), citations ('[Source: file, Slide X]'), and honest uncertainty ('I don't know'). Without it, the LLM hallucinates."],
        ["7", "Different queries need different strategies", "Comparison queries need two separate retrievals. One-size-fits-all retrieval does not work for enterprise use cases."],
        ["8", "Start with one model, measure, then optimize", "GPT-4.1 Mini handles 80% of queries well. Only upgrade to Full for complex reasoning. Never optimize before measuring."],
        ["9", "Evaluation is not optional", "Without metrics, improvement is guesswork. With metrics (retrieval hit, faithfulness, citation), you know exactly what to fix."],
        ["10", "Build modularly, migrate incrementally", "Every module can be replaced independently. ChromaDB to Azure AI Search? Change 2 files. Preprocessor improvement? Change 1 file. The rest stays the same."],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# TOPIC 22: APPENDIX
# ═══════════════════════════════════════════════════════════════

doc.add_heading('22. Appendix', level=1)

doc.add_heading('22.1 Tech Stack', level=2)
add_table(doc,
    ["Component", "Technology", "Version", "Purpose"],
    [
        ["Language", "Python", "3.11+", "Core development"],
        ["Embeddings", "Azure OpenAI text-embedding-3-small", "2024-12-01-preview", "Text to 1536-dim vectors"],
        ["LLM", "Azure OpenAI GPT-4.1", "2024-12-01-preview", "Answer generation"],
        ["Vector DB", "ChromaDB", "0.5+", "Local vector storage and search"],
        ["API Framework", "FastAPI", "0.111+", "REST API"],
        ["API Server", "Uvicorn", "0.30+", "ASGI server for FastAPI"],
        ["UI", "Streamlit", "1.35+", "Chat web interface"],
        ["PPT Extraction", "python-pptx", "0.6+", "PowerPoint text extraction"],
        ["PDF Extraction", "PyMuPDF (fitz)", "1.24+", "PDF text extraction"],
        ["DOCX Extraction", "python-docx", "1.1+", "Word document extraction"],
        ["Configuration", "Pydantic Settings", "2.3+", "Typed config from .env"],
        ["Testing", "Pytest", "8.2+", "Unit and integration tests"],
    ]
)

doc.add_heading('22.2 API Endpoints', level=2)
add_table(doc,
    ["Method", "Endpoint", "Request Body", "Response", "Purpose"],
    [
        ["GET", "/health", "None", "{status, service, version}", "Health check for monitoring"],
        ["POST", "/ingest", "None", "{files_processed, chunks_created, errors, time}", "Trigger full document ingestion"],
        ["POST", "/query", "{question: str, client?: str}", "{answer, sources, intent, tokens, time}", "Ask a question, get cited answer"],
        ["GET", "/stats", "None", "{total_chunks, clients, files, section_types}", "Index statistics"],
        ["GET", "/clients", "None", "[client_names]", "List available clients"],
    ]
)

doc.add_heading('22.3 Project Folder Structure', level=2)
add_code(doc,
    'arch-knowledge-assistant/\n'
    '|\n'
    '+-- config/\n'
    '|   +-- settings.py              # Central configuration\n'
    '|\n'
    '+-- data/\n'
    '|   +-- documents/               # Client PPTs go here\n'
    '|   |   +-- DataStories/\n'
    '|   |   +-- TechNova/\n'
    '|   |   +-- MediFlow/\n'
    '|   |   +-- RetailEdge/\n'
    '|   +-- chromadb/                # Vector store (auto-created)\n'
    '|\n'
    '+-- src/\n'
    '|   +-- ingestion/\n'
    '|   |   +-- scanner.py           # Discover documents\n'
    '|   |   +-- extractor.py         # PPT/PDF/DOCX to text\n'
    '|   |   +-- preprocessor.py      # Clean + classify\n'
    '|   |   +-- chunker.py           # Merge/split + metadata\n'
    '|   +-- indexing/\n'
    '|   |   +-- indexer.py           # Embed + store in ChromaDB\n'
    '|   +-- retrieval/\n'
    '|   |   +-- retriever.py         # Semantic search + filtering\n'
    '|   +-- generation/\n'
    '|   |   +-- generator.py         # LLM answers + citations\n'
    '|   +-- pipeline/\n'
    '|   |   +-- ingest_pipeline.py   # E2E ingestion\n'
    '|   |   +-- query_pipeline.py    # E2E query\n'
    '|   +-- api/\n'
    '|       +-- main.py              # FastAPI REST API\n'
    '|\n'
    '+-- ui/\n'
    '|   +-- app.py                   # Streamlit chat UI\n'
    '|\n'
    '+-- eval/\n'
    '|   +-- golden_qa.json           # Test Q&A pairs\n'
    '|   +-- evaluate.py              # Evaluation script\n'
    '|\n'
    '+-- tests/\n'
    '|   +-- test_ingestion.py\n'
    '|   +-- test_retrieval.py\n'
    '|   +-- test_generation.py\n'
    '|\n'
    '+-- docs/                        # This guide!\n'
    '+-- results/                     # Evaluation outputs\n'
    '+-- .env                         # Secrets (never in git)\n'
    '+-- .env.template                # Config template\n'
    '+-- .gitignore\n'
    '+-- requirements.txt\n'
    '+-- README.md')

doc.add_heading('22.4 Key Configuration Parameters', level=2)
add_table(doc,
    ["Parameter", "Default Value", "What It Controls", "When to Change"],
    [
        ["CHUNK_MIN_LENGTH", "100 chars", "Minimum chunk size (below = merge)", "If too many tiny chunks appear"],
        ["CHUNK_MAX_LENGTH", "1500 chars", "Maximum chunk size (above = split)", "If retrieval is too broad or too narrow"],
        ["CHUNK_OVERLAP", "150 chars", "Overlap when splitting large chunks", "If context is lost at chunk boundaries"],
        ["RETRIEVAL_TOP_K", "8", "Max chunks per query", "If answers lack detail (increase) or have noise (decrease)"],
        ["RETRIEVAL_MIN_RELEVANCE", "0.35", "Minimum similarity score", "If irrelevant chunks appear (increase) or good chunks are dropped (decrease)"],
        ["GENERATION_MAX_TOKENS", "2000", "Max answer length", "If answers are truncated (increase)"],
        ["GENERATION_TEMPERATURE", "0.1", "LLM creativity (0=deterministic, 1=creative)", "Keep low for factual RAG; only increase for brainstorming"],
    ]
)

doc.add_heading('22.5 Glossary', level=2)
add_table(doc,
    ["Term", "Definition"],
    [
        ["RAG", "Retrieval-Augmented Generation — pattern where LLM answers using retrieved documents"],
        ["Embedding", "A vector (list of numbers) that captures the meaning of text"],
        ["Vector Database", "Database optimized for storing and searching embedding vectors"],
        ["Chunk", "A piece of text stored as one unit in the vector database"],
        ["Token", "A sub-word unit that AI models process (1 token ~ 4 characters)"],
        ["Cosine Similarity", "Measure of angle between two vectors (1 = identical, 0 = unrelated)"],
        ["Top-K", "Number of most similar chunks to retrieve per query"],
        ["Grounding", "Ensuring LLM answers are based on provided documents, not training data"],
        ["Hallucination", "When an LLM generates plausible but incorrect information"],
        ["System Prompt", "Instructions that define how the LLM should behave"],
        ["Context Window", "Maximum number of tokens an LLM can process in one call"],
        ["BM25", "Classic keyword-based search algorithm (used in hybrid search)"],
        ["Reranking", "Second-pass scoring of retrieved results using a cross-encoder model"],
        ["Idempotent", "Operation that produces the same result regardless of how many times it runs"],
        ["RBAC", "Role-Based Access Control — users see only data they are authorized for"],
        ["OCR", "Optical Character Recognition — extracting text from images"],
    ]
)

# ── Final page ──
doc.add_page_break()

doc.add_heading('End of Document', level=1)
add_body(doc, '')
p = doc.add_paragraph()
run = p.add_run(
    'This guide was created as part of the Architecture Knowledge Assistant project.\n\n'
    'The project demonstrates that building a production-grade RAG system is not about '
    'using the latest framework or the most powerful model — it is about understanding '
    'each layer deeply, making intentional design decisions, and engineering for real-world messiness.\n\n'
    'Every concept in this guide was learned by building, breaking, and fixing a real system '
    'with real enterprise documents.\n\n'
    'The best way to learn RAG is to build RAG.'
)
run.font.italic = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(80, 80, 80)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
p2 = doc.add_paragraph('Architecture Team — June 2026')
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in p2.runs:
    run.font.size = Pt(14)
    run.bold = True

# ═══════════════════════════════════════════════════════════════
# SAVE
# ═══════════════════════════════════════════════════════════════

output_path = "docs/RAG_Guide_Batch5.docx"
doc.save(output_path)
print(f"✅ Saved: {output_path}")
print(f"   Topics covered: 17-22 (FINAL BATCH)")
print(f"   (Modernization, Gaps, Advanced Patterns, Roadmap, Takeaways, Appendix)")
print(f"")
print(f"📚 ALL 5 BATCHES COMPLETE!")